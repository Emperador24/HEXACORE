#!/usr/bin/env python3
"""Genera la sección de requisitos funcionales del SRS a partir del .xlsx.

La fuente de verdad de los casos de uso es `Submission/CU_eventos_completo.xlsx`:
32 hojas, una por CU, con la misma plantilla. Copiar esas fichas a mano dentro
del LaTeX las condenaría a divergir del .xlsx a la primera corrección, así que
la sección se genera.

    python3 generar-requisitos-cu.py      # reescribe RequisitosCasosUso.tex

`EspecificacionRequisitos.tex` lo incluye con \\input. No editar el .tex
generado: se pierde en la siguiente corrida.

Sin dependencias más allá de openpyxl (pip install openpyxl).
"""

from pathlib import Path
import re
import unicodedata

import openpyxl

AQUI = Path(__file__).resolve().parent
LIBRO = AQUI.parent / 'Submission' / 'CU_eventos_completo.xlsx'
SALIDA = AQUI / 'RequisitosCasosUso.tex'

# Los responsables salen del campo "Autor:" de cada hoja; esto solo normaliza
# las variantes de escritura que aparecen en el libro.
RESPONSABLES = {
    'daniel cristancho': 'Daniel Cristancho',
    'samuel emperador': 'Samuel Emperador',
    'sebastian sanchez': 'Sebastián Sánchez',
    'diego coronado': 'Diego Coronado',
}

# Etiqueta de la fila en el .xlsx -> nombre corto con el que se usa aquí.
CAMPOS = {
    'objetivo en contexto': 'objetivo',
    'actores participantes': 'actores',
    'requisitos asociados (codigos)': 'asociados',
    'entradas': 'entradas',
    'pre-condiciones (escenario de exito)': 'precondiciones',
    'salidas': 'salidas',
    'post-condiciones (escenario de exito)': 'postcondiciones',
    'flujos o caminos alternativos (dentro del escenario de exito)': 'alternativos',
    'caminos de excepcion (camino que no es normal o esperado)': 'excepciones',
    'extensiones': 'extensiones',
    'atributos de calidad asociados': 'calidad',
    'infraestructura no trivial utilizada': 'infraestructura',
}


def sin_tildes(texto: str) -> str:
    """Para comparar etiquetas sin depender de tildes ni mayúsculas."""
    plano = unicodedata.normalize('NFD', texto)
    return ''.join(c for c in plano if unicodedata.category(c) != 'Mn').lower()


def escapar(texto: str) -> str:
    """Texto plano -> LaTeX. El contenido viene de una hoja de cálculo, así que
    puede traer cualquier carácter especial."""
    reemplazos = {
        '\\': r'\textbackslash{}', '&': r'\&', '%': r'\%', '$': r'\$',
        '#': r'\#', '_': r'\_', '{': r'\{', '}': r'\}',
        '~': r'\textasciitilde{}', '^': r'\textasciicircum{}',
    }
    return ''.join(reemplazos.get(c, c) for c in texto)


def lineas(valor) -> list[str]:
    """Una celda multilínea -> lista de líneas limpias."""
    if valor is None:
        return []
    return [l.strip() for l in str(valor).replace('\r', '').split('\n') if l.strip()]


def primera_celda(fila) -> str:
    for celda in fila[1:]:
        if celda not in (None, ''):
            return str(celda)
    return ''


def leer_hoja(hoja):
    """Una hoja del libro -> diccionario con los campos de la ficha."""
    datos = {'id': hoja.title, 'nombre': '', 'autor': '', 'pasos': []}
    en_flujo = False

    for fila in hoja.iter_rows(values_only=True):
        etiqueta = sin_tildes(str(fila[0] or '').replace('\n', ' ')).strip()
        etiqueta = re.sub(r'\s+', ' ', etiqueta).rstrip(':')

        if etiqueta == 'autor':
            datos['autor'] = primera_celda(fila)
            continue

        if etiqueta == 'id caso de uso':
            # La misma fila lleva "Nombre:" en la columna C y el valor en la D.
            valores = [str(c) for c in fila[1:] if c not in (None, '')]
            if len(valores) >= 3:
                datos['nombre'] = valores[2]
            continue

        if etiqueta.startswith('flujo basico de exito'):
            en_flujo = True
            continue

        if etiqueta in CAMPOS:
            en_flujo = False
            datos[CAMPOS[etiqueta]] = lineas(primera_celda(fila))
            continue

        if en_flujo:
            # El flujo ocupa cinco columnas y alterna una fila por participante:
            #   A = nº del paso del actor    B = lo que hace el actor
            #   D = nº del paso del sistema  E = lo que hace el sistema
            # (C va vacía, separa visualmente las dos mitades).
            if etiqueta == 'no. secuencia':
                continue
            col = list(fila) + [None] * 5
            if col[1]:
                datos['pasos'].append(('Actor', str(col[0] or '').strip(), str(col[1]).strip()))
            if col[4]:
                datos['pasos'].append(('Sistema', str(col[3] or '').strip(), str(col[4]).strip()))

    return datos


def bloque_lista(titulo: str, valores: list[str]) -> str:
    """Un campo de la ficha como fila de descripción."""
    if not valores:
        return ''
    if len(valores) == 1:
        return f'\\item[{escapar(titulo)}] {escapar(valores[0])}\n'
    puntos = '\n'.join(f'  \\item {escapar(v)}' for v in valores)
    return (f'\\item[{escapar(titulo)}] \\hfill\n'
            f'\\begin{{itemize}}[leftmargin=1.2em, topsep=1pt]\n{puntos}\n\\end{{itemize}}\n')


def ficha(datos: dict) -> str:
    cu = escapar(datos['id'])
    nombre = escapar(datos['nombre'] or datos['id'])
    autor = RESPONSABLES.get(sin_tildes(datos['autor']).strip(), datos['autor'])

    partes = [
        f'\\subsection{{{cu} --- {nombre}}}',
        f'\\label{{cu:{datos["id"].lower()}}}',
        f'\\textit{{Responsable: {escapar(autor)}.}}',
        '',
        '\\begin{description}[leftmargin=!, labelwidth=3.4cm, style=nextline, topsep=3pt]',
    ]

    for titulo, clave in (
        ('Objetivo', 'objetivo'),
        ('Actores', 'actores'),
        ('CU relacionados', 'asociados'),
        ('Entradas', 'entradas'),
        ('Pre-condiciones', 'precondiciones'),
        ('Salidas', 'salidas'),
        ('Post-condiciones', 'postcondiciones'),
    ):
        partes.append(bloque_lista(titulo, datos.get(clave, [])))

    partes.append('\\end{description}')

    if datos['pasos']:
        partes += [
            '',
            '\\paragraph{Flujo básico de éxito}',
            '\\begin{longtable}{P{1.1cm}P{1.8cm}P{10.5cm}}',
            '\\toprule',
            '\\textbf{Paso} & \\textbf{Quién} & \\textbf{Acción}\\\\',
            '\\midrule',
            '\\endhead',
        ]
        for i, (quien, numero, texto) in enumerate(datos['pasos'], start=1):
            paso = escapar(numero) if numero else str(i)
            partes.append(f'{paso} & {quien} & {escapar(texto)}\\\\')
        partes += ['\\bottomrule', '\\end{longtable}']

    for titulo, clave in (
        ('Flujos alternativos', 'alternativos'),
        ('Caminos de excepción', 'excepciones'),
        ('Extensiones', 'extensiones'),
        ('Atributos de calidad asociados', 'calidad'),
        ('Infraestructura no trivial', 'infraestructura'),
    ):
        valores = datos.get(clave, [])
        if not valores:
            continue
        partes += ['', f'\\paragraph{{{titulo}}}']
        if len(valores) == 1:
            partes.append(escapar(valores[0]))
        else:
            partes.append('\\begin{itemize}[leftmargin=1.4em, topsep=1pt]')
            partes += [f'  \\item {escapar(v)}' for v in valores]
            partes.append('\\end{itemize}')

    return '\n'.join(p for p in partes if p is not None)


def main() -> None:
    libro = openpyxl.load_workbook(LIBRO, data_only=True)
    fichas = [leer_hoja(libro[nombre]) for nombre in libro.sheetnames]

    encabezado = (
        '% ARCHIVO GENERADO por generar-requisitos-cu.py — no editar a mano.\n'
        f'% Fuente: Submission/CU_eventos_completo.xlsx ({len(fichas)} casos de uso).\n\n'
    )
    SALIDA.write_text(encabezado + '\n\n'.join(ficha(f) for f in fichas) + '\n', encoding='utf-8')

    por_autor: dict[str, list[str]] = {}
    for f in fichas:
        autor = RESPONSABLES.get(sin_tildes(f['autor']).strip(), f['autor'] or 'sin autor')
        por_autor.setdefault(autor, []).append(f['id'])

    print(f'{SALIDA.name}: {len(fichas)} casos de uso')
    for autor, cus in sorted(por_autor.items()):
        print(f'  {autor}: {len(cus)} — {", ".join(cus)}')


if __name__ == '__main__':
    main()
